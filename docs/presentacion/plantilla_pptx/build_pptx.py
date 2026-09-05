#!/usr/bin/env python3
"""Genera la plantilla WatchGate en .pptx (para subir a Google Slides).

Replica la identidad visual de la presentación LaTeX/Beamer:
- paleta navy / cyan
- portada con foto + velo navy, logo nuevo en blanco (sin logos UMA)
- títulos en versalita con regla cian
- pie de página discreto
- diapositivas-plantilla de ejemplo (contenido, diagrama, tarjetas, tabla)
"""
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

ASSETS = Path(__file__).parent.parent           # docs/presentacion/
LOGOS = ASSETS / "logos"
OUT = ASSETS / "WatchGate_plantilla.pptx"

# ---------- paleta (idéntica a main.tex) ----------
NAVY   = RGBColor(0x00, 0x29, 0x59)
NAVY2  = RGBColor(0x12, 0x3A, 0x66)
CYAN   = RGBColor(0x00, 0xAD, 0xCB)
PAPER  = RGBColor(0xF5, 0xF6, 0xF8)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
INK    = RGBColor(0x14, 0x20, 0x2E)
INKSOFT= RGBColor(0x4B, 0x58, 0x6A)
LINE   = RGBColor(0xD6, 0xDC, 0xE5)
BOXFILL= RGBColor(0xE7, 0xEE, 0xF6)
RED    = RGBColor(0xB2, 0x3A, 0x2E)
AMBER  = RGBColor(0x9C, 0x6B, 0x0C)
GREEN  = RGBColor(0x2E, 0x7D, 0x4F)

F_HEAD = "Lato"        # titulares (sustituye a Optima) — nativa en Google Slides
F_BODY = "PT Serif"    # cuerpo (sustituye a Charter) — nativa en Google Slides
F_MONO = "Roboto Mono" # código / datos (sustituye a Menlo) — nativa en Google Slides

EMU_W, EMU_H = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width = EMU_W
prs.slide_height = EMU_H
BLANK = prs.slide_layouts[6]


# ---------- helpers ----------
def _set_alpha(fill_elem, pct):
    """pct = opacidad 0-100 sobre el color sólido de un fill."""
    srgb = fill_elem.find(qn("a:srgbClr"))
    a = srgb.makeelement(qn("a:alpha"), {"val": str(int(pct * 1000))})
    srgb.append(a)


def _no_shadow(shp):
    """Elimina cualquier sombra heredada del tema (empty effectLst)."""
    spPr = shp._element.spPr
    for tag in ("a:effectLst", "a:effectDag"):
        for el in spPr.findall(qn(tag)):
            spPr.remove(el)
    spPr.append(spPr.makeelement(qn("a:effectLst"), {}))


def add_slide():
    return prs.slides.add_slide(BLANK)


def bg(slide, color):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, EMU_W, EMU_H)
    r.fill.solid(); r.fill.fore_color.rgb = color
    r.line.fill.background()
    r.shadow.inherit = False; _no_shadow(r)
    slide.shapes._spTree.remove(r._element)
    slide.shapes._spTree.insert(2, r._element)
    return r


def rect(slide, x, y, w, h, color, line_color=None, line_w=0.75, alpha=None, rounded=False):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, x, y, w, h)
    if rounded:
        try:
            shp.adjustments[0] = 0.045
        except Exception:
            pass
    shp.fill.solid(); shp.fill.fore_color.rgb = color
    if alpha is not None:
        _set_alpha(shp.fill.fore_color._xFill.find(qn("a:solidFill")) if False else
                   shp.fill._xPr.find(qn("a:solidFill")), alpha)
    if line_color is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line_color
        shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    _no_shadow(shp)
    return shp


def tb(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
       space_after=6, line_spacing=1.06, wrap=True):
    """runs: lista de párrafos; cada párrafo es lista de (texto, dict-estilo)."""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        for text, st in para:
            r = p.add_run(); r.text = text
            r.font.name = st.get("font", F_BODY)
            r.font.size = Pt(st.get("size", 16))
            r.font.bold = st.get("bold", False)
            r.font.italic = st.get("italic", False)
            r.font.color.rgb = st.get("color", INK)
            if st.get("spacing"):  # tracking en pt
                rPr = r._r.get_or_add_rPr()
                rPr.set("spc", str(int(st["spacing"] * 100)))
    return box


def frametitle(slide, label):
    """Título de diapositiva: versalita navy + regla cian (como en main.tex)."""
    tb(slide, Inches(0.62), Inches(0.42), Inches(12.1), Inches(0.7),
       [[(label.upper(), dict(font=F_HEAD, size=25, bold=True, color=NAVY, spacing=0.6))]])
    ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.62), Inches(1.16),
                                Inches(12.1), Pt(2.2))
    ln.fill.solid(); ln.fill.fore_color.rgb = CYAN
    ln.line.fill.background(); ln.shadow.inherit = False; _no_shadow(ln)


def footer(slide, n):
    tb(slide, Inches(0.62), Inches(7.03), Inches(12.1), Inches(0.32),
       [[("WATCHGATE", dict(font=F_MONO, size=8, color=INKSOFT, spacing=0.4)),
         ("      Cátedra de Ciberseguridad · UMA · 2026", dict(font=F_MONO, size=8, color=INKSOFT))]])
    tb(slide, Inches(11.7), Inches(7.03), Inches(1.03), Inches(0.32),
       [[(str(n), dict(font=F_MONO, size=8, color=INKSOFT))]], align=PP_ALIGN.RIGHT)


def lead(slide, text):
    tb(slide, Inches(0.62), Inches(1.4), Inches(12.1), Inches(0.6),
       [[(text, dict(font=F_BODY, size=19, color=INK))]])


def body(slide, text, y=2.05, size=14):
    tb(slide, Inches(0.62), Inches(y), Inches(12.1), Inches(1.2),
       [[(text, dict(font=F_BODY, size=size, color=INKSOFT))]], line_spacing=1.18)


# ======================================================================
# 1 · PORTADA
# ======================================================================
s = add_slide()
pic = s.shapes.add_picture(str(ASSETS / "malaga.jpeg"), 0, 0, EMU_W, EMU_H)
velo = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, EMU_W, EMU_H)
velo.fill.solid(); velo.fill.fore_color.rgb = NAVY
_set_alpha(velo.fill._xPr.find(qn("a:solidFill")), 82)
velo.line.fill.background(); velo.shadow.inherit = False; _no_shadow(velo)

s.shapes.add_picture(str(LOGOS / "watchgate_white.png"), Inches(0.95), Inches(0.85),
                     height=Inches(1.65))

tb(s, Inches(0.98), Inches(2.95), Inches(11.4), Inches(0.5),
   [[("CÁTEDRA DE CIBERSEGURIDAD · UNIVERSIDAD DE MÁLAGA",
      dict(font=F_HEAD, size=15, bold=True, color=CYAN, spacing=1.0))]])

tb(s, Inches(0.93), Inches(3.5), Inches(11.4), Inches(1.5),
   [[("Watch", dict(font=F_HEAD, size=66, bold=True, color=WHITE)),
     ("Gate", dict(font=F_HEAD, size=66, bold=True, color=CYAN))]])

tb(s, Inches(0.98), Inches(4.95), Inches(10.8), Inches(1.0),
   [[("Sistema de ", dict(font=F_BODY, size=18, color=WHITE)),
     ("scoring", dict(font=F_BODY, size=18, italic=True, color=WHITE)),
     (" de riesgo para la revisión automatizada de ", dict(font=F_BODY, size=18, color=WHITE)),
     ("pull requests", dict(font=F_BODY, size=18, italic=True, color=WHITE)),
     (" en pipelines CI/CD", dict(font=F_BODY, size=18, color=WHITE))]],
   line_spacing=1.25)

tb(s, Inches(0.98), Inches(6.35), Inches(11.4), Inches(0.8),
   [[("Javier Martín Jurado · Pablo Jiménez Castro · Pablo Ayllón García",
      dict(font=F_BODY, size=12, color=WHITE))],
    [("Málaga, septiembre de 2026", dict(font=F_BODY, size=12, color=RGBColor(0xC5,0xD0,0xDE))) ]],
   space_after=3)

# ======================================================================
# 2 · ÍNDICE
# ======================================================================
s = add_slide(); bg(s, PAPER)
frametitle(s, "Índice")
col_items = [
    ["01  El problema", "02  Estado del arte", "03  Arquitectura general",
     "04  El contrato entre capas", "05  Las cinco capas"],
    ["06  Capa estática", "07  Capa de dependencias", "08  Capa de vulnerabilidades",
     "09  Capa de reputación", "10  Capa semántica (opcional)", "11  Flujo con agentes de IA (MCP)"],
]
for c, items in enumerate(col_items):
    x = Inches(0.9 + c * 6.2)
    runs = []
    for it in items:
        num, rest = it.split("  ", 1)
        runs.append([(num + "   ", dict(font=F_HEAD, size=17, bold=True, color=CYAN)),
                     (rest, dict(font=F_BODY, size=17, color=INK))])
    tb(s, x, Inches(1.75), Inches(5.7), Inches(4.8), runs, space_after=16, line_spacing=1.1)
footer(s, 2)

# ======================================================================
# 3 · PLANTILLA — diapositiva de contenido
# ======================================================================
s = add_slide(); bg(s, PAPER)
frametitle(s, "Título de la sección")
lead(s, "Frase-tesis en una línea: la idea que quieres que se lleven de esta diapositiva")
body(s, "Párrafo de apoyo en gris, 2–3 líneas como máximo. Aquí va el matiz, el dato "
        "concreto o el ejemplo que sostiene la frase de arriba. Evita listas largas: "
        "una idea por diapositiva.")
# viñetas de ejemplo
bullets = [
    "Viñeta con el marcador cian rectangular, como en la memoria",
    "Segunda idea, al mismo nivel",
    "Sub-idea o cita textual, en cursiva",
]
by = 3.5
for i, bt in enumerate(bullets):
    mk = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.66), Inches(by + i*0.62 + 0.11),
                            Inches(0.16), Pt(4.5))
    mk.fill.solid(); mk.fill.fore_color.rgb = CYAN if i < 2 else LINE
    mk.line.fill.background(); mk.shadow.inherit = False; _no_shadow(mk)
    tb(s, Inches(1.0), Inches(by + i*0.62), Inches(11.6), Inches(0.55),
       [[(bt, dict(font=F_BODY, size=15, color=INK, italic=(i == 2)))]])
footer(s, 3)

# ======================================================================
# 4 · PLANTILLA — diagrama a pantalla
# ======================================================================
s = add_slide(); bg(s, PAPER)
frametitle(s, "Arquitectura general")
tb(s, Inches(0.62), Inches(1.32), Inches(12.1), Inches(0.5),
   [[("Un núcleo agnóstico de plataforma · cuatro capas deterministas · la capa LLM es ", dict(font=F_BODY, size=15, color=INK)),
     ("opcional", dict(font=F_BODY, size=15, bold=True, color=NAVY2))]])
_dh = Inches(5.4)                       # alto disponible bajo el título
with Image.open(ASSETS / "diagrama_capas.png") as _im:
    _dw = Emu(int(_dh * _im.size[0] / _im.size[1]))   # ancho según proporción real
s.shapes.add_picture(str(ASSETS / "diagrama_capas.png"),
                     Emu(int((EMU_W - _dw) / 2)), Inches(1.58), width=_dw, height=_dh)
footer(s, 4)

# ======================================================================
# 5 · PLANTILLA — diagrama a pantalla (agentes de IA + servidor MCP)
# ======================================================================
s = add_slide(); bg(s, PAPER)
frametitle(s, "Flujo con agentes de IA (MCP)")
tb(s, Inches(0.62), Inches(1.32), Inches(12.1), Inches(0.5),
   [[("El agente llama al servidor MCP ", dict(font=F_BODY, size=15, color=INK)),
     ("antes de abrir el PR", dict(font=F_BODY, size=15, bold=True, color=NAVY2)),
     (" — mismo núcleo de WatchGate, shift-left", dict(font=F_BODY, size=15, color=INK))]])
_dh2 = Inches(5.4)
with Image.open(ASSETS / "diagrama_mcp_agentes.png") as _im2:
    _dw2 = Emu(int(_dh2 * _im2.size[0] / _im2.size[1]))
s.shapes.add_picture(str(ASSETS / "diagrama_mcp_agentes.png"),
                     Emu(int((EMU_W - _dw2) / 2)), Inches(1.58), width=_dw2, height=_dh2)
footer(s, 5)

# ======================================================================
# 6 · PLANTILLA — tres tarjetas
# ======================================================================
s = add_slide(); bg(s, PAPER)
frametitle(s, "Tres columnas / tarjetas")
lead(s, "Para comparativas de 3 elementos: casos reales, alternativas, fases…")
cards = [
    ("2024 · CVE-2024-3094", "XZ Utils", "Puerta trasera en liblzma introducida por "
     "un colaborador de confianza; el payload vivía en ficheros de test y scripts de build."),
    ("2026 · AUR", "Atomic Arch", "Más de 1.500 paquetes comprometidos falsificando "
     "nombre y correo de commits anteriores para simular continuidad."),
    ("2026 · Wiz Research", "prt-scan", "Seis cuentas generando PRs maliciosos a escala "
     "con IA, con el payload adaptado al lenguaje de cada repositorio."),
]
cw, gap = Inches(3.95), Inches(0.28)
for i, (tag, title, txt) in enumerate(cards):
    x = Inches(0.62) + i * (cw + gap)
    card = rect(s, x, Inches(2.5), cw, Inches(2.85), WHITE, line_color=NAVY2, line_w=0.75, rounded=True)
    tb(s, x + Inches(0.24), Inches(2.75), cw - Inches(0.48), Inches(2.5),
       [[(tag, dict(font=F_MONO, size=9, color=CYAN))],
        [(title, dict(font=F_HEAD, size=17, bold=True, color=NAVY2))],
        [(txt, dict(font=F_BODY, size=11.5, color=INKSOFT))]],
       space_after=8, line_spacing=1.16)
footer(s, 6)

# ======================================================================
# 7 · PLANTILLA — tabla
# ======================================================================
s = add_slide(); bg(s, PAPER)
frametitle(s, "Tabla comparativa")
lead(s, "Cabecera navy, filas alternas, la fila de WatchGate resaltada")
rows = [
    ("Herramienta", "Abierta", "Multiplat.", "Identidad vs. hist.", "Explicable"),
    ("Socket.dev / GuardDog", "~", "—", "—", "~"),
    ("BewAIre (Datadog, interno)", "—", "—", "—", "—"),
    ("Phylum / Aikido / Snyk", "—", "—", "~", "n/a"),
    ("WatchGate", "sí", "sí", "sí", "sí"),
]
tbl_shape = s.shapes.add_table(len(rows), 5, Inches(1.1), Inches(2.2),
                               Inches(11.1), Inches(3.6))
table = tbl_shape.table
table.columns[0].width = Inches(3.9)
for c in range(1, 5):
    table.columns[c].width = Inches(1.8)
for ri, row in enumerate(rows):
    for ci, val in enumerate(row):
        cell = table.cell(ri, ci)
        cell.margin_left = Inches(0.12); cell.margin_right = Inches(0.08)
        cell.margin_top = Inches(0.06); cell.margin_bottom = Inches(0.06)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        if ri == 0:
            cell.fill.solid(); cell.fill.fore_color.rgb = NAVY
            col, bold, size = WHITE, True, 12
        elif row[0] == "WatchGate":
            cell.fill.solid(); cell.fill.fore_color.rgb = BOXFILL
            col, bold, size = NAVY2, True, 12
        else:
            cell.fill.solid(); cell.fill.fore_color.rgb = WHITE if ri % 2 else PAPER
            col, bold, size = INK, False, 12
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.CENTER
        r = p.add_run(); r.text = val
        r.font.name = F_BODY if ci == 0 else F_MONO
        r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = col
footer(s, 7)

# ======================================================================
# 8 · PLANTILLA — cierre
# ======================================================================
s = add_slide()
bg(s, NAVY)
s.shapes.add_picture(str(LOGOS / "watchgate_white.png"), Inches(5.9), Inches(1.7),
                     height=Inches(1.7))
tb(s, Inches(1.0), Inches(3.7), Inches(11.3), Inches(1.0),
   [[("Watch", dict(font=F_HEAD, size=44, bold=True, color=WHITE)),
     ("Gate", dict(font=F_HEAD, size=44, bold=True, color=CYAN))]], align=PP_ALIGN.CENTER)
tb(s, Inches(1.0), Inches(4.8), Inches(11.3), Inches(0.6),
   [[("¿Preguntas?", dict(font=F_BODY, size=18, color=RGBColor(0xC5,0xD0,0xDE)))]],
   align=PP_ALIGN.CENTER)
tb(s, Inches(1.0), Inches(6.6), Inches(11.3), Inches(0.5),
   [[("github.com/pablojimz/watch_gate", dict(font=F_MONO, size=11, color=CYAN))]],
   align=PP_ALIGN.CENTER)

prs.save(str(OUT))
print("OK ->", OUT)
